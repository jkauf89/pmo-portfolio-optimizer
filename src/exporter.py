"""Excel export module for PMO portfolio optimization results."""

from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_portfolio_results(
    df: pd.DataFrame,
    budget_cap: float,
    capacity_cap: int,
    output_path: str = "output/portfolio_results.xlsx",
) -> Path:
    """Exports optimization results into an executive-styled, multi-tab Excel workbook."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    funded_df = df[df["selected"]].copy()
    cut_df = df[~df["selected"]].copy()

    total_initiatives = len(df)
    funded_count = len(funded_df)
    cut_count = len(cut_df)

    total_cost_funded = funded_df["cost_k_usd"].sum()
    total_weeks_funded = funded_df["eng_weeks"].sum()
    total_value_captured = funded_df["adjusted_score"].sum()

    budget_pct = (total_cost_funded / budget_cap) if budget_cap > 0 else 0
    capacity_pct = (total_weeks_funded / capacity_cap) if capacity_cap > 0 else 0

    # 1. Structure the Executive Scorecard with Categories
    summary_data = [
        {"Category": "Portfolio Scope", "Metric": "Total Initiatives Evaluated", "Value": str(total_initiatives)},
        {"Category": "Portfolio Scope", "Metric": "Initiatives Funded", "Value": f"{funded_count} ({funded_count / total_initiatives:.0%})"},
        {"Category": "Portfolio Scope", "Metric": "Initiatives Deferred / Cut", "Value": str(cut_count)},
        {"Category": "Financial Capital", "Metric": "Total Budget Cap", "Value": f"${budget_cap:,.1f}k"},
        {"Category": "Financial Capital", "Metric": "Budget Committed", "Value": f"${total_cost_funded:,.1f}k"},
        {"Category": "Financial Capital", "Metric": "Budget Utilization Rate", "Value": f"{budget_pct:.1%}"},
        {"Category": "Capacity & Resourcing", "Metric": "Total Capacity Cap", "Value": f"{capacity_cap} person-weeks"},
        {"Category": "Capacity & Resourcing", "Metric": "Capacity Allocated", "Value": f"{total_weeks_funded} person-weeks"},
        {"Category": "Capacity & Resourcing", "Metric": "Capacity Utilization Rate", "Value": f"{capacity_pct:.1%}"},
        {"Category": "Strategic ROI", "Metric": "Total Strategic Value Captured", "Value": f"{total_value_captured:.2f} pts"},
    ]
    summary_df = pd.DataFrame(summary_data)

    # 2. Write Sheets
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Executive Summary", index=False)
        funded_df.to_excel(writer, sheet_name="Funded Initiatives", index=False)
        cut_df.to_excel(writer, sheet_name="Cut & Deferred", index=False)
        df.to_excel(writer, sheet_name="Full Master List", index=False)

        # 3. Apply Professional Styling
        wb = writer.book

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # Navy
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=11)
        category_font = Font(name="Calibri", size=11, bold=True, color="1F497D")

        thin_side = Side(style="thin", color="D9D9D9")
        grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        for ws in wb.worksheets:
            # Enable gridlines in Excel view
            ws.views.sheetView[0].showGridLines = True

            # Style Header Row
            ws.row_dimensions[1].height = 24
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Style Data Rows
            for row in ws.iter_rows(min_row=2):
                ws.row_dimensions[row[0].row].height = 20
                for cell in row:
                    cell.font = data_font
                    cell.border = grid_border
                    # If on Executive Summary and Column 1 (Category), style as bold section
                    if ws.title == "Executive Summary" and cell.column == 1:
                        cell.font = category_font
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    elif isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-calculate and assign readable column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    return path